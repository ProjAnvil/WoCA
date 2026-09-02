#!/usr/bin/env python3
"""woca office 插件的内嵌脚本（由 woca 物化后在 uv 管理的 venv 中执行）。

协议：子命令 + 参数，stdout 输出单个 JSON 对象；失败时 {"ok": false, "error": ...}
且退出码为 1。woca 的 OfficeRead / DocxWrite / XlsxWrite / PptxWrite 工具是唯一调用方。

用法：
  office.py read <file> [--max-chars N]
  office.py docx-write --output <file> --spec '<json>'
  office.py xlsx-write --output <file> --spec '<json>'
  office.py pptx-write --output <file> --spec '<json>'

spec 均为 JSON 字符串（由工具层 marshal 后经 argv 传入，不落临时文件，
与 pdf 插件一样保持沙箱执行器兼容）：

  docx: {"title": "...", "blocks": [
           {"type": "heading", "text": "...", "level": 1},   # level 0 = Title 样式
           {"type": "paragraph", "text": "..."},
           {"type": "bullets", "items": ["...", "..."]},
           {"type": "table", "header": ["A", "B"], "rows": [["1", "2"]]},
           {"type": "pagebreak"} ]}
  xlsx: {"sheets": [{"name": "Sheet1", "rows": [["Name", "Qty"], ["apple", "3"]],
                     "header_row": true}]}
         单元格为字符串；数字样式的串自动转数值，"=" 开头是公式，"'" 前缀强制文本。
  pptx: {"title": "...", "subtitle": "...",             # 可选封面
         "slides": [{"title": "...",
                     "bullets": [{"text": "...", "level": 0}],
                     "notes": "..."}]}
"""

from __future__ import annotations

import argparse
import datetime as _dt
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pptx import Presentation

# 病态输入的护栏：超出即报错，避免 agent 一次生成把内存/文件拖爆。
MAX_BLOCKS = 5000
MAX_ITEMS = 2000
MAX_TABLE_ROWS = 20000
MAX_TABLE_COLS = 100
MAX_SHEETS = 50
MAX_XLSX_ROWS = 100_000
MAX_XLSX_COLS = 1000
MAX_XLSX_CELLS = 200_000
MAX_SLIDES = 500
MAX_BULLETS = 200
READ_MAX_COLS = 200
READ_MAX_ROWS_PER_SHEET = 1000


def out(obj: dict) -> None:
    sys.stdout.write(json.dumps(obj, ensure_ascii=False))
    sys.stdout.write("\n")
    sys.stdout.flush()


def fail(message: str) -> None:
    out({"ok": False, "error": message})
    sys.exit(1)


# ---------------------------------------------------------------------------
# 读取：三种格式 → Markdown 行流


def _md_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return str(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    return str(value).replace("|", "\\|").replace("\n", " ")


def docx_lines(path: str):
    doc = Document(path)
    num_counter = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            style = para.style.name if para.style is not None else ""
            if not text:
                if not style.startswith("List Number"):
                    num_counter = 0
                continue
            if style == "Title":
                num_counter = 0
                yield "# " + text
                continue
            m = re.fullmatch(r"Heading (\d)", style)
            if m:
                num_counter = 0
                yield "#" * min(int(m.group(1)), 9) + " " + text
                continue
            if style.startswith("List Bullet"):
                num_counter = 0
                depth = 0
                m = re.fullmatch(r"List Bullet (\d)", style)
                if m:
                    depth = max(0, int(m.group(1)) - 1)
                yield "  " * depth + "- " + text
                continue
            if style.startswith("List Number"):
                num_counter += 1
                depth = 0
                m = re.fullmatch(r"List Number (\d)", style)
                if m:
                    depth = max(0, int(m.group(1)) - 1)
                yield "  " * depth + f"{num_counter}. " + text
                continue
            num_counter = 0
            yield text
        elif child.tag == qn("w:tbl"):
            num_counter = 0
            table = Table(child, doc)
            rows = [[_md_cell(c.text) for c in row.cells] for row in table.rows]
            yield from _md_table_lines(rows)


def _md_table_lines(rows: list[list[str]]):
    if not rows:
        return
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    yield "| " + " | ".join(rows[0]) + " |"
    yield "|" + "|".join([" --- "] * width) + "|"
    for r in rows[1:]:
        yield "| " + " | ".join(r) + " |"


def xlsx_lines(path: str):
    wb = load_workbook(filename=path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            yield f"## 工作表：{ws.title}"
            emitted = 0
            header_done = False
            width = 0
            for raw in ws.iter_rows(values_only=True):
                if raw is None:
                    continue
                if emitted >= READ_MAX_ROWS_PER_SHEET:
                    yield f"（已截断：每工作表最多展示 {READ_MAX_ROWS_PER_SHEET} 行）"
                    break
                if all(c is None for c in raw):
                    continue
                row = list(raw[:READ_MAX_COLS])
                if len(raw) > READ_MAX_COLS:
                    row.append(f"（+{len(raw) - READ_MAX_COLS} 列未展示）")
                if not header_done:
                    width = len(row)
                    yield "| " + " | ".join(_md_cell(c) for c in row) + " |"
                    yield "|" + "|".join([" --- "] * width) + "|"
                    header_done = True
                else:
                    if len(row) < width:
                        row += [None] * (width - len(row))
                    yield "| " + " | ".join(_md_cell(c) for c in row) + " |"
                emitted += 1
            if not header_done:
                yield "（空工作表）"
            yield ""
    finally:
        wb.close()


def pptx_lines(path: str):
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = title_shape.text.strip()
        head = f"## 幻灯片 {i}"
        if title:
            head += f"：{title}"
        yield head
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_table:
                rows = [[_md_cell(c.text) for c in row.cells] for row in shape.table.rows]
                yield from _md_table_lines(rows)
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        yield "  " * min(para.level, 4) + "- " + text
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                yield "> 备注：" + notes.replace("\n", " ")
        yield ""


def cmd_read(args: argparse.Namespace) -> dict:
    path = Path(args.file)
    if not path.is_file():
        raise ValueError(f"文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix in (".docx", ".docm"):
        fmt, lines = "docx", docx_lines(str(path))
    elif suffix in (".xlsx", ".xlsm"):
        fmt, lines = "xlsx", xlsx_lines(str(path))
    elif suffix in (".pptx", ".pptm"):
        fmt, lines = "pptx", pptx_lines(str(path))
    else:
        raise ValueError(f"不支持的格式 {suffix!r}；支持 .docx / .xlsx / .pptx（含 .docm / .xlsm / .pptm 宏变体）")
    limit = args.max_chars if args.max_chars > 0 else 20000
    chunks: list[str] = []
    used = 0
    truncated = False
    for piece in lines:
        piece = piece.rstrip()
        remain = limit - used
        if remain <= 0:
            truncated = True
            break
        if len(piece) > remain:
            piece = piece[:remain]
            truncated = True
        chunks.append(piece)
        used += len(piece) + 1
    return {
        "ok": True,
        "file": args.file,
        "format": fmt,
        "truncated": truncated,
        "markdown": "\n".join(chunks).strip(),
    }


# ---------------------------------------------------------------------------
# 写入：结构化 spec → 文件


def _require(cond: bool, message: str) -> None:
    if not cond:
        raise ValueError(message)


def _as_str(value: object, where: str) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)) or value is None:
        return _md_cell(value)
    raise ValueError(f"{where} 必须是字符串，得到 {type(value).__name__}")


def _parse_spec(raw: str) -> dict:
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"spec 不是合法 JSON：{exc}") from exc
    _require(isinstance(spec, dict), "spec 必须是 JSON 对象")
    return spec


def _write_out(path: str) -> Path:
    target = Path(path)
    if target.parent and not target.parent.exists():
        target.parent.mkdir(parents=True, exist_ok=True)
    return target


def cmd_docx_write(args: argparse.Namespace) -> dict:
    spec = _parse_spec(args.spec)
    title = spec.get("title") or ""
    blocks = spec.get("blocks") or []
    _require(isinstance(title, str), "title 必须是字符串")
    _require(isinstance(blocks, list), "blocks 必须是数组")
    _require(bool(title) or blocks, "需要 title 或至少一个 block")
    _require(len(blocks) <= MAX_BLOCKS, f"blocks 超过上限 {MAX_BLOCKS}")

    doc = Document()
    if title:
        doc.add_heading(title, 0)
    for i, raw in enumerate(blocks):
        where = f"blocks[{i}]"
        _require(isinstance(raw, dict), f"{where} 必须是对象")
        kind = raw.get("type")
        if kind == "heading":
            text = _as_str(raw.get("text", ""), where)
            _require(bool(text), f"{where} heading 需要 text")
            level = raw.get("level", 1)
            _require(isinstance(level, int), f"{where} level 必须是整数")
            doc.add_heading(text, max(0, min(level, 9)))
        elif kind == "paragraph":
            text = _as_str(raw.get("text", ""), where)
            _require(bool(text), f"{where} paragraph 需要 text")
            doc.add_paragraph(text)
        elif kind == "bullets":
            items = raw.get("items") or []
            _require(isinstance(items, list), f"{where} items 必须是数组")
            _require(0 < len(items) <= MAX_ITEMS, f"{where} items 数量需在 1..{MAX_ITEMS}")
            for item in items:
                doc.add_paragraph(_as_str(item, where), style="List Bullet")
        elif kind == "table":
            header = raw.get("header") or []
            rows = raw.get("rows") or []
            _require(isinstance(header, list) and 0 < len(header) <= MAX_TABLE_COLS,
                     f"{where} header 需要 1..{MAX_TABLE_COLS} 列")
            _require(isinstance(rows, list) and len(rows) <= MAX_TABLE_ROWS,
                     f"{where} rows 数量需在 0..{MAX_TABLE_ROWS}")
            cells = [[_as_str(c, where) for c in row] for row in rows]
            table = doc.add_table(rows=1 + len(cells), cols=len(header))
            table.style = "Table Grid"
            for col, text in enumerate(header):
                cell = table.rows[0].cells[col]
                cell.text = _as_str(text, where)
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.bold = True
            for r, row in enumerate(cells, start=1):
                for col in range(len(header)):
                    table.rows[r].cells[col].text = row[col] if col < len(row) else ""
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(f"{where} 未知 type {kind!r}（heading/paragraph/bullets/table/pagebreak）")

    target = _write_out(args.output)
    doc.save(str(target))
    return {"ok": True, "output": args.output, "blocks": len(blocks), "bytes": target.stat().st_size}


_NUMBER_RE = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d+)?$")


def _coerce_cell(text: str):
    """字符串 → 单元格值：数字样式转数值（保留前导零为文本），"'" 前缀强制文本。"""
    if text.startswith("'"):
        return text[1:]
    if text.startswith("="):
        return text
    if _NUMBER_RE.fullmatch(text):
        return float(text) if "." in text else int(text)
    return text


def _display_width(text: str) -> int:
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


_SHEET_BAD = set(':\\/?*[]')


def cmd_xlsx_write(args: argparse.Namespace) -> dict:
    spec = _parse_spec(args.spec)
    sheets = spec.get("sheets") or []
    _require(isinstance(sheets, list) and 0 < len(sheets) <= MAX_SHEETS,
             f"sheets 需要 1..{MAX_SHEETS} 个")
    wb = Workbook()
    total_cells = 0
    total_rows = 0
    seen_names: set[str] = set()
    for i, raw in enumerate(sheets):
        where = f"sheets[{i}]"
        _require(isinstance(raw, dict), f"{where} 必须是对象")
        rows = raw.get("rows") or []
        _require(isinstance(rows, list) and 0 < len(rows) <= MAX_XLSX_ROWS,
                 f"{where} rows 需要 1..{MAX_XLSX_ROWS} 行")
        name = _as_str(raw.get("name") or f"Sheet{i + 1}", where)
        name = "".join("_" if ch in _SHEET_BAD else ch for ch in name)[:31]
        base = name or f"Sheet{i + 1}"
        unique, n = base, 2
        while unique in seen_names:
            suffix = f"~{n}"
            unique = base[: 31 - len(suffix)] + suffix
            n += 1
        seen_names.add(unique)

        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = unique
        header_row = bool(raw.get("header_row"))
        widths: list[int] = []
        for r, row in enumerate(rows, start=1):
            _require(isinstance(row, list), f"{where} rows[{r - 1}] 必须是数组")
            _require(len(row) <= MAX_XLSX_COLS, f"{where} rows[{r - 1}] 超过 {MAX_XLSX_COLS} 列")
            total_cells += len(row)
            _require(total_cells <= MAX_XLSX_CELLS, f"总单元格超过上限 {MAX_XLSX_CELLS}")
            for c, value in enumerate(row, start=1):
                text = _as_str(value, where)
                cell_value = _coerce_cell(text)
                ws.cell(row=r, column=c, value=cell_value)
                if header_row and r == 1:
                    ws.cell(row=r, column=c).font = Font(bold=True)
                display = text if text.startswith(("=", "'")) else _md_cell(cell_value)
                while len(widths) < c:
                    widths.append(0)
                widths[c - 1] = max(widths[c - 1], min(_display_width(display), 60))
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = max(8, w + 2)
        if header_row and len(rows) > 1:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{len(rows)}"
        total_rows += len(rows)

    target = _write_out(args.output)
    wb.save(str(target))
    return {
        "ok": True,
        "output": args.output,
        "sheets": len(sheets),
        "rows": total_rows,
        "bytes": target.stat().st_size,
    }


def cmd_pptx_write(args: argparse.Namespace) -> dict:
    spec = _parse_spec(args.spec)
    cover_title = _as_str(spec.get("title") or "", "title")
    cover_subtitle = _as_str(spec.get("subtitle") or "", "subtitle")
    slides = spec.get("slides") or []
    _require(isinstance(slides, list) and len(slides) <= MAX_SLIDES,
             f"slides 数量需在 0..{MAX_SLIDES}")
    _require(bool(cover_title) or bool(slides), "需要封面 title 或至少一张 slide")

    prs = Presentation()
    slides_added = 0
    if cover_title:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = cover_title
        subtitle = slide.placeholders[1]
        subtitle.text = cover_subtitle
        slides_added += 1
    for i, raw in enumerate(slides):
        where = f"slides[{i}]"
        _require(isinstance(raw, dict), f"{where} 必须是对象")
        title = _as_str(raw.get("title") or "", where)
        bullets = raw.get("bullets") or []
        _require(isinstance(bullets, list) and len(bullets) <= MAX_BULLETS,
                 f"{where} bullets 数量需在 0..{MAX_BULLETS}")
        _require(bool(title) or bool(bullets), f"{where} 需要 title 或 bullets")
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        if title:
            slide.shapes.title.text = title
        if bullets:
            body = slide.placeholders[1].text_frame
            first = True
            for b in bullets:
                if isinstance(b, str):
                    text, level = b, 0
                elif isinstance(b, dict):
                    text = _as_str(b.get("text", ""), where)
                    level = b.get("level", 0)
                    _require(isinstance(level, int), f"{where} bullet level 必须是整数")
                else:
                    raise ValueError(f"{where} bullet 必须是字符串或 {{text, level}}")
                _require(bool(text), f"{where} bullet 需要 text")
                para = body.paragraphs[0] if first else body.add_paragraph()
                first = False
                para.text = text
                para.level = max(0, min(level, 4))
        notes = _as_str(raw.get("notes") or "", where)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        slides_added += 1

    target = _write_out(args.output)
    prs.save(str(target))
    return {"ok": True, "output": args.output, "slides": slides_added, "bytes": target.stat().st_size}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="office.py", description="woca 内嵌办公文档处理脚本")
    sub = parser.add_subparsers(dest="op", required=True)

    p = sub.add_parser("read")
    p.add_argument("file")
    p.add_argument("--max-chars", type=int, default=20000)

    for name in ("docx-write", "xlsx-write", "pptx-write"):
        p = sub.add_parser(name)
        p.add_argument("--output", required=True)
        p.add_argument("--spec", required=True)
    return parser


HANDLERS = {
    "read": cmd_read,
    "docx-write": cmd_docx_write,
    "xlsx-write": cmd_xlsx_write,
    "pptx-write": cmd_pptx_write,
}


def main() -> None:
    args = build_parser().parse_args()
    try:
        result = HANDLERS[args.op](args)
    except FileNotFoundError as exc:
        fail(f"文件不存在：{exc}")
    except ValueError as exc:
        fail(str(exc))
    except Exception as exc:  # noqa: BLE001 - 顶层兜底，错误内容回传给 agent
        fail(f"{type(exc).__name__}: {exc}")
    out(result)


if __name__ == "__main__":
    main()
